from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from lark_channel import (  # type: ignore[import-untyped]
    Conversation,
    FileContent,
    Identity,
    InboundMessage,
    ResourceDescriptor,
)

from feishu_dify_gateway.autodev_content import (
    RequirementContentRejected,
    attachment_from_message,
    normalize_text,
)


class FakeChannel:
    def __init__(self, payload: bytes) -> None:
        self.payload = payload

    async def download_resource(self, file_key: str, **_: object) -> bytes:
        return self.payload


def file_message(name: str) -> InboundMessage:
    return InboundMessage(
        id="file-message",
        create_time=0,
        conversation=Conversation(chat_id="oc-chat", chat_type="p2p"),
        sender=Identity(open_id="ou-owner"),
        content=FileContent(file_key="file-key", file_name=name),
        resources=[ResourceDescriptor(type="file", file_key="file-key", file_name=name)],
    )


@pytest.mark.asyncio
async def test_txt_and_docx_attachments_are_extracted_with_digest(tmp_path: Path) -> None:
    txt = await attachment_from_message(
        FakeChannel(b"Title\n\nA requirement."),
        file_message("requirement.txt"),
        max_bytes=1024,
        max_chars=1000,
    )
    assert txt.source_kind == "txt"
    assert txt.text == "Title\n\nA requirement."
    assert len(txt.content_sha256) == 64

    document = Document()
    document.add_heading("DOCX requirement", level=1)
    document.add_paragraph("Use the bounded workflow.")
    stream = BytesIO()
    document.save(stream)
    docx = await attachment_from_message(
        FakeChannel(stream.getvalue()),
        file_message("requirement.docx"),
        max_bytes=1024 * 1024,
        max_chars=1000,
    )
    assert "Use the bounded workflow." in docx.text


@pytest.mark.asyncio
async def test_unsupported_and_oversized_files_fail_closed() -> None:
    with pytest.raises(RequirementContentRejected, match="仅支持"):
        await attachment_from_message(
            FakeChannel(b"data"),
            file_message("requirement.png"),
            max_bytes=1024,
            max_chars=1000,
        )
    with pytest.raises(RequirementContentRejected, match="大小"):
        await attachment_from_message(
            FakeChannel(b"too large"),
            file_message("requirement.txt"),
            max_bytes=2,
            max_chars=1000,
        )


def test_text_normalization_is_bounded() -> None:
    normalized = normalize_text(
        "  first  \r\nsecond  ",
        title="",
        max_chars=100,
        source_kind="text",
    )
    assert normalized.text == "first\nsecond"
    with pytest.raises(RequirementContentRejected, match="长度"):
        normalize_text("12345", title="", max_chars=4, source_kind="text")
