from __future__ import annotations

import hashlib
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader


class RequirementContentRejected(ValueError):
    """The provider content is unsupported or violates an intake limit."""


@dataclass(frozen=True, slots=True)
class NormalizedRequirement:
    title: str
    text: str
    content_sha256: str
    source_kind: str


def normalize_text(
    text: str,
    *,
    title: str,
    max_chars: int,
    source_kind: str,
) -> NormalizedRequirement:
    normalized = "\n".join(line.rstrip() for line in text.replace("\r\n", "\n").split("\n")).strip()
    if not normalized:
        raise RequirementContentRejected("需求正文为空，请发送文本或可提取文本的文件。")
    if len(normalized) > max_chars:
        raise RequirementContentRejected("需求正文超过当前版本允许的最大长度。")
    safe_title = title.strip() or normalized.splitlines()[0][:80]
    return NormalizedRequirement(
        title=safe_title[:512],
        text=normalized,
        content_sha256=hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
        source_kind=source_kind,
    )


def text_from_message(message: Any, *, max_chars: int) -> NormalizedRequirement | None:
    content = getattr(message, "content", None)
    kind = getattr(content, "kind", "")
    if kind == "text":
        return normalize_text(
            str(getattr(content, "text", "")),
            title="",
            max_chars=max_chars,
            source_kind="text",
        )
    if kind == "post":
        return normalize_text(
            str(getattr(content, "text", "")),
            title=str(getattr(content, "title", "")),
            max_chars=max_chars,
            source_kind="post",
        )
    content_text = getattr(message, "content_text", "")
    if (
        isinstance(content_text, str)
        and content_text.strip()
        and not getattr(message, "resources", None)
    ):
        return normalize_text(
            content_text,
            title="",
            max_chars=max_chars,
            source_kind="text",
        )
    return None


async def attachment_from_message(
    channel: Any,
    message: Any,
    *,
    max_bytes: int,
    max_chars: int,
) -> NormalizedRequirement:
    resources = list(getattr(message, "resources", ()) or ())
    content = getattr(message, "content", None)
    file_key = str(getattr(content, "file_key", ""))
    file_name = str(getattr(content, "file_name", "") or "")
    resource = next(
        (
            item
            for item in resources
            if not file_key or str(getattr(item, "file_key", "")) == file_key
        ),
        None,
    )
    if resource is None and resources:
        resource = resources[0]
    if resource is None:
        raise RequirementContentRejected("未找到可下载的需求附件。")
    resource_type = str(getattr(resource, "type", ""))
    if resource_type != "file":
        raise RequirementContentRejected("当前只支持文本类文件附件。")
    file_key = str(getattr(resource, "file_key", "")) or file_key
    file_name = file_name or str(getattr(resource, "file_name", "") or "")
    suffix = Path(file_name).suffix.lower()
    allowed = {".txt", ".md", ".docx", ".pdf"}
    if suffix not in allowed:
        raise RequirementContentRejected("当前仅支持 .txt、.md、.docx 和可提取文本的 .pdf。")
    if not file_key:
        raise RequirementContentRejected("需求附件缺少受控资源标识。")
    try:
        raw = await channel.download_resource(
            file_key,
            resource_type="file",
            message_id=str(getattr(message, "id", "")) or None,
        )
    except Exception as exc:
        raise RequirementContentRejected("需求附件下载失败，请重新发送文件。") from exc
    if not isinstance(raw, bytes) or not raw:
        raise RequirementContentRejected("需求附件为空或下载失败。")
    if len(raw) > max_bytes:
        raise RequirementContentRejected("需求附件超过当前版本允许的大小。")
    try:
        with tempfile.TemporaryDirectory(prefix="autodev-requirement-") as directory:
            path = Path(directory) / Path(file_name).name
            path.write_bytes(raw)
            text = _extract(path, suffix)
    except RequirementContentRejected:
        raise
    except Exception as exc:
        raise RequirementContentRejected(
            "需求文件无法提取文本，请发送文本、Markdown 或 DOCX。"
        ) from exc
    return normalize_text(
        text,
        title=Path(file_name).stem,
        max_chars=max_chars,
        source_kind=suffix[1:],
    )


def _extract(path: Path, suffix: str) -> str:
    if suffix in {".txt", ".md"}:
        try:
            return path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as exc:
            raise RequirementContentRejected("文本附件必须使用 UTF-8 编码。") from exc
    if suffix == ".docx":
        document = Document(str(path))
        chunks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            chunks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(chunks)
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise RequirementContentRejected("当前不支持加密 PDF，请发送文本或 DOCX。")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if not text.strip():
            raise RequirementContentRejected("当前不支持扫描版或无文本 PDF，请发送文本或 DOCX。")
        return text
    raise RequirementContentRejected("不支持的需求文件类型。")
